import re
import docx
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def md_to_docx(md_path, docx_path):
    doc = docx.Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    
    for line in lines:
        line = line.strip('\n')
        
        # Handle code blocks (mermaid)
        if line.startswith('```'):
            in_code_block = not in_code_block
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style = 'Normal'
            # try to use courier if possible or just normal text
            for run in p.runs:
                run.font.name = 'Courier New'
            continue
            
        if not line.strip():
            continue
            
        # Headings
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=4)
        elif line.startswith('- '):
            # List item
            p = doc.add_paragraph(style='List Bullet')
            _add_formatted_text(p, line[2:])
        elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or line.startswith('4. ') or line.startswith('5. ') or line.startswith('6. ') or line.startswith('7. ') or line.startswith('8. ') or line.startswith('9. '):
            # Ordered list item
            p = doc.add_paragraph(style='List Number')
            _add_formatted_text(p, re.sub(r'^\d+\.\s', '', line))
        else:
            p = doc.add_paragraph()
            _add_formatted_text(p, line)
            
    doc.save(docx_path)

def _add_formatted_text(paragraph, text):
    # Very basic bold and italic parser
    # splits by **
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            # Handle italics
            subparts = re.split(r'(\*.*?\*)', part)
            for subpart in subparts:
                if subpart.startswith('*') and subpart.endswith('*') and not subpart.startswith('**'):
                    run = paragraph.add_run(subpart[1:-1])
                    run.italic = True
                else:
                    paragraph.add_run(subpart)

if __name__ == '__main__':
    md_to_docx(r'C:\Users\DELL\.gemini\antigravity\brain\53612b62-9e2e-412f-af48-87e3f9ecf094\project_chapters_1_to_5.md', r'C:\Users\DELL\.gemini\antigravity\brain\53612b62-9e2e-412f-af48-87e3f9ecf094\Telemed_Project_Chapters_1_to_5.docx')
